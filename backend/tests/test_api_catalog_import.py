"""Каталоги: вендоры с ответственными, анкеты LMS, откат битых строк импорта."""

from io import BytesIO

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from tests.conftest import ADMIN, MANAGER

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def book(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def run(client: AsyncClient, rows: list[list[object]], import_type: str) -> dict:
    preview = await client.post(
        "/api/v1/imports",
        files={"file": ("data.xlsx", book(rows), XLSX_TYPE)},
        data={"type": import_type},
        headers=ADMIN,
    )
    assert preview.status_code == 201, preview.text
    run_id = preview.json()["run"]["id"]
    result = await client.post(f"/api/v1/imports/{run_id}/commit", json={}, headers=ADMIN)
    assert result.status_code == 200, result.text
    return result.json()


VENDORS = [
    ["Компания", "Продукт", "ФИО", "Телефон", "Почта", "Способ связи"],
    ["ООО «ТДата»", "«RT.DataLake», «RT.Warehouse»", "Смирнова Анна Петровна",
     "+7 (911) 222-33-44", "smirnova.ap@example.ru", "Чат в ТГ"],
    ["ООО «РТК ИТ Плюс»", "«AKOLA»", "Попова Мария Владимировна", "+7 (933) 444-55-66",
     "popova.mv@example.ru", "Чат в ТГ"],
    ["ООО «РТК ИТ Плюс»", "«Яга»", "Соколов Алексей Андреевич", "+7 (944) 555-66-77",
     "sokolov.aa@example.ru", "Чат в ТГ"],
]  # fmt: skip


async def test_vendor_catalog_in_case_format(client: AsyncClient) -> None:
    result = await run(client, VENDORS, "vendors")
    assert result["run"]["rows_created"] == 2  # две компании
    assert result["run"]["rows_updated"] == 1  # вторая строка той же компании
    assert result["run"]["rows_failed"] == 0

    vendors = {
        item["name"]: item
        for item in (await client.get("/api/v1/catalog/vendors", headers=MANAGER)).json()
    }
    assert {contact["full_name"] for contact in vendors["ООО «РТК ИТ Плюс»"]["contacts"]} == {
        "Попова Мария Владимировна",
        "Соколов Алексей Андреевич",
    }
    tdata = vendors["ООО «ТДата»"]
    assert tdata["contacts"][0]["contact_channel"] == "Чат в ТГ"

    # Два продукта из одной ячейки, кавычки-ёлочки сняты, ответственный - у продукта.
    products = {
        item["name"]: item
        for item in (await client.get("/api/v1/catalog/products", headers=MANAGER)).json()
    }
    assert {"RT.DataLake", "RT.Warehouse", "AKOLA", "Яга"} <= set(products)
    assert products["RT.Warehouse"]["contact_id"] == tdata["contacts"][0]["id"]
    assert products["AKOLA"]["vendor_id"] == vendors["ООО «РТК ИТ Плюс»"]["id"]


LMS_HEADERS = [
    "Фамилия", "Имя", "Отчествопри наличии)", "Номер телефона", "Email", "СНИЛС",
    "Серия паспорта", "Номер паспорта", "Пол", "Образование", "Регион регистрации",
]  # fmt: skip


async def test_lms_questionnaire_import_keeps_only_needed_fields(client: AsyncClient) -> None:
    preview = await client.post(
        "/api/v1/imports",
        files={
            "file": (
                "lms.xlsx",
                book(
                    [
                        LMS_HEADERS,
                        [
                            "Гусева",
                            "Анна",
                            "Петровна",
                            79000234365,
                            "guseva.a@example.com",
                            "000-000-000 00",
                            "0000",
                            "000000",
                            "Ж",
                            "Высшее образование – бакалавриат",
                            "Москва",
                        ],
                        [
                            "Тарасов",
                            "Олег",
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            "М",
                            None,
                            None,
                        ],
                        [
                            "Власова",
                            "Ирина",
                            None,
                            79004583434,
                            None,
                            None,
                            None,
                            None,
                            "X",
                            None,
                            None,
                        ],
                    ]  # fmt: skip
                ),
                XLSX_TYPE,
            )
        },
        data={"type": "learners"},
        headers=ADMIN,
    )
    body = preview.json()
    # Испорченный заголовок «Отчествопри наличии)» узнан, паспорт и СНИЛС -
    # не сопоставлены ни с одним полем, то есть не будут загружены.
    assert body["suggested_mapping"]["middle_name"] == "Отчествопри наличии)"
    assert "СНИЛС" not in body["suggested_mapping"].values()

    result = (
        await client.post(
            f"/api/v1/imports/{body['run']['id']}/commit", json={}, headers=ADMIN
        )
    ).json()
    assert result["run"]["rows_created"] == 1
    # Без телефона и почты анкету не с чем сопоставить; пол «X» не распознан.
    assert result["run"]["rows_failed"] == 2
    messages = " ".join(error["message"] for error in result["errors"])
    assert "нужен телефон или почта" in messages
    assert "М или Ж" in messages


async def test_contacts_import_needs_known_university(
    client: AsyncClient, university: dict
) -> None:
    result = await run(
        client,
        [
            ["Название ВУЗа", "ФИО", "Должность", "Почта"],
            [university["name"], "Гусева Анна Петровна", "Проректор", "guseva@example.edu"],
            ["Неизвестный вуз", "Тарасов Олег", None, None],
        ],
        "contacts",
    )
    assert result["run"]["rows_created"] == 1
    assert result["run"]["rows_failed"] == 1
    assert "нет в справочнике" in result["errors"][0]["message"]

    detail = (
        await client.get(f"/api/v1/universities/{university['id']}", headers=MANAGER)
    ).json()
    assert detail["contacts"][0]["position"] == "Проректор"


CATALOG_HEADERS = [
    "Название ВУЗа", "Вендор", "ПО", "Номер договора", "Подписание лицензии",
    "Срок действия лицензии (год)", "Статус по передачи", "ФИО Менеджера",
    "Ответственные от ВУЗа", "Комментарий",
]  # fmt: skip


async def test_bad_row_leaves_nothing_behind(client: AsyncClient) -> None:
    """Строка с битой датой отклоняется целиком: ни вуза, ни договора от неё."""
    result = await run(
        client,
        [
            CATALOG_HEADERS,
            ["Хороший вуз", None, None, "ДГ-1", None, None, None, None, "Гусева Анна", None],
            [
                "Плохой вуз",
                "Вендор",
                "Продукт",
                "ДГ-2",
                "31.02.2026",
                None,
                None,
                None,
                None,
                None,
            ],
        ],  # fmt: skip
        "catalog",
    )
    assert (result["run"]["rows_created"], result["run"]["rows_failed"]) == (1, 1)
    assert result["errors"][0]["row_number"] == 3  # номер строки в файле

    universities = (await client.get("/api/v1/universities", headers=ADMIN)).json()
    assert [item["name"] for item in universities["items"]] == ["Хороший вуз"]

    # Ответственный от вуза назначен на договор, а не просто заведён у вуза.
    contracts = (await client.get("/api/v1/contracts", headers=ADMIN)).json()
    detail = (
        await client.get(f"/api/v1/contracts/{contracts['items'][0]['id']}", headers=ADMIN)
    ).json()
    assert detail["contacts"][0]["contact"]["full_name"] == "Гусева Анна"
    assert detail["contacts"][0]["role"] == "Ответственный от вуза"


async def test_import_is_not_committed_twice(client: AsyncClient) -> None:
    preview = await client.post(
        "/api/v1/imports",
        files={
            "file": (
                "catalog.xlsx",
                book([CATALOG_HEADERS, ["Вуз", *[None] * 2, "ДГ-7"]]),
                XLSX_TYPE,
            )
        },
        data={"type": "catalog"},
        headers=ADMIN,
    )
    run_id = preview.json()["run"]["id"]
    first = await client.post(f"/api/v1/imports/{run_id}/commit", json={}, headers=ADMIN)
    assert first.status_code == 200
    second = await client.post(f"/api/v1/imports/{run_id}/commit", json={}, headers=ADMIN)
    assert second.status_code == 409


async def test_template_has_headers_only(client: AsyncClient) -> None:
    response = await client.get("/api/v1/imports/template?type=vendors", headers=ADMIN)
    sheet = load_workbook(BytesIO(response.content)).worksheets[0]
    rows = [row for row in sheet.iter_rows(values_only=True) if any(row)]
    assert rows == [("Компания", "Продукт", "ФИО", "Телефон", "Почта", "Способ связи")]
    assert "Обязательное" in sheet["A1"].comment.text


async def test_catalog_items_are_edited_and_deactivated(client: AsyncClient) -> None:
    program = (
        await client.post("/api/v1/catalog/programs", json={"name": "Курс"}, headers=ADMIN)
    ).json()
    duplicate = await client.post(
        "/api/v1/catalog/programs", json={"name": "курс"}, headers=ADMIN
    )
    assert duplicate.status_code == 409

    updated = await client.patch(
        f"/api/v1/catalog/programs/{program['id']}",
        json={"name": "Курс 2.0", "is_active": False},
        headers=ADMIN,
    )
    assert updated.json()["name"] == "Курс 2.0"
    assert updated.json()["is_active"] is False

    denied = await client.patch(
        f"/api/v1/catalog/programs/{program['id']}", json={"name": "x"}, headers=MANAGER
    )
    assert denied.status_code == 403
