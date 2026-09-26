"""Файлы для ручной проверки загрузок: импорт каталогов и вложения.

Запуск:
    python -m scripts.testdata ../testdata      # из каталога backend
    make testdata                               # то же в контейнере разработки

Кладёт в каталог:

    import/        таблицы для импорта каталогов (требование 1 ТЗ) и файлы
                   в форматах кейсодержателя: «Вендоры», анкета LMS
    integrations/  ответы сайта и LMS в формате кейсодержателя - для загрузки
                   файлом в разделе «Интеграции»
    attachments/   по файлу каждого формата вложений из ТЗ и один недопустимый

Что в каждом файле и чего ждать от загрузки - в testdata/README.md.
Люди во всех файлах вымышленные.
"""

import argparse
import json
from collections.abc import Callable
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from scripts.demo import files, learning, sheets, vendors

# файл -> (название листа, таблица, формат)
IMPORTS: dict[str, tuple[str, Callable[[], sheets.Sheet], str]] = {
    "01-catalog.xlsx": ("Сводный каталог", sheets.catalog, "xlsx"),
    "02-catalog.xls": ("Сводный каталог", sheets.catalog, "xls"),
    "03-catalog-other-headers.xlsx": ("Каталог", sheets.catalog_other_headers, "xlsx"),
    "04-catalog-errors.xlsx": ("Каталог с ошибками", sheets.catalog_errors, "xlsx"),
    "05-universities.xlsx": ("Вузы", sheets.universities, "xlsx"),
    "06-programs.xlsx": ("ИТ-программы", sheets.programs, "xlsx"),
    "07-products.xlsx": ("ИТ-продукты", sheets.products, "xlsx"),
    "08-vendors.xlsx": ("Лист1", vendors.sheet, "xlsx"),
}

# Шаблон анкеты LMS кейсодержателя: заголовки как в исходном файле, вместе
# с испорченными при выгрузке («Отчествопри наличии)»).
LMS_HEADERS = [
    "Фамилия", "Имя", "Отчествопри наличии)", "Номер телефона", "Email", "СНИЛС",
    "Серия паспорта", "Номер паспорта", "Кем выдан паспорт", "Дата выдачи паспорта",
    "Код подразделения", "Пол", "Дата рождения", "Регион регистрации",
    "Населенный пункт регистрации", "Улица регистрации", "Дом регистрации",
    "Квартира регистрации", "Индекс регистрации", "Имядательный падеж)",
    "Фамилиядательный падеж)", "Отчестводательный падеж)", "Образование",
    "Профессия по диплому", "Учебное заведение по диплому", "Фамилия, указанная в дипломе",
    "Номер диплома", "Серия диплома", "Регистрационный номер диплома", "Дата выдачи диплома",
]  # fmt: skip

# Заявки сайта и анкеты LMS для файлов: те же вымышленные люди.
_PEOPLE = [
    # Фамилия, Имя, Отчество, Курс, поток, пол, образование, регион, дошёл ли до LMS
    ("Гусева", "Анна", "Петровна", "Анализ данных без программирования", 1, "Ж",
     "Высшее образование – бакалавриат", "Москва", True),
    ("Тарасов", "Олег", "Юрьевич", "Инженер-тестировщик", 1, "М",
     "Среднее профессиональное образование", "Московская область", True),
    ("Власова", "Ирина", "Николаевна",
     "Управление ИТ-проектами на базе программного продукта ПАО «Ростелеком»", 2, "Ж",
     "Высшее образование – специалитет, магистратура", "Санкт-Петербург", True),
    ("Орехов", "Дмитрий", "Сергеевич", "Промпт-инжиниринг", 3, "М",
     "Высшее образование – бакалавриат", "Республика Татарстан", False),
    ("Сафина", "Алина", "Рашидовна", "Python-разработчик с использованием инструментов ИИ",
     4, "Ж", "Среднее общее образование - 11 классов", "Новосибирская область", True),
]  # fmt: skip
_NUMBERS = (
    "ORD-20260313051569-OYJRVN",  # испорченная секунда: даты в номере нет
    "ORD-202605130654453-GDJIKG",
    "ORD-20261721184559-AJIJEN",  # испорченный месяц
    "ORD-20260522061330-2LT0MG",
    "ORD-20260904075403-ZXFSZX",
)
_EDUCATION_LEVELS = (
    "Без образования",
    "Основное общее образование - 9 классов",
    "Среднее общее образование - 11 классов",
    "Среднее профессиональное образование",
    "Высшее образование – бакалавриат",
    "Высшее образование – специалитет, магистратура",
    "Высшее образование – подготовка кадров высшей квалификации",
)


def _person_contacts(index: int, last: str) -> tuple[str, int, str]:
    phone = f"7 (900) 0{10 + index}-{20 + index}-{30 + index}"
    digits = int("".join(char for char in phone if char.isdigit()))
    email = f"{learning.translit(last)}.{index}@example.com"
    return phone, digits, email


def site_applications() -> list[object]:
    """Ответ сайта в формате кейсодержателя: пустой первый элемент - как в оригинале."""
    items: list[object] = [None]
    for index, (last, first, middle, course, stream, *_rest) in enumerate(_PEOPLE):
        phone, _, email = _person_contacts(index, last)
        items.append(
            {
                "Номер заявки": _NUMBERS[index],
                "Курс": course,
                "Фамилия": last,
                "Имя": first,
                "Отчество": middle,
                "Телефон": phone,
                "Email": email,
                "Номер потока": stream,
            }
        )
    return items


def lms_learners() -> list[dict[str, object]]:
    """Анкеты LMS для JSON: все поля шаблона, документы - нулями."""
    rows = []
    for index, person in enumerate(_PEOPLE):
        last, first, middle, _course, _stream, gender, education, region, enrolled = person
        if not enrolled:
            continue
        _, digits, email = _person_contacts(index, last)
        row: dict[str, object] = dict.fromkeys(LMS_HEADERS, None)
        row |= {
            "Фамилия": last, "Имя": first, "Отчествопри наличии)": middle,
            "Номер телефона": digits, "Email": email, "СНИЛС": "000-000-000 00",
            "Серия паспорта": "0000", "Номер паспорта": "000000", "Пол": gender,
            "Образование": education, "Регион регистрации": region,
        }  # fmt: skip
        rows.append(row)
    return rows


def lms_workbook() -> bytes:
    """Анкета LMS как в файле кейсодержателя: лист данных и лист справочников."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append(LMS_HEADERS)
    for row in lms_learners():
        sheet.append([row[header] for header in LMS_HEADERS])
    lists = workbook.create_sheet("Лист2")
    for index, level in enumerate(_EDUCATION_LEVELS, start=1):
        lists.cell(row=index, column=2, value=level)
    lists.cell(row=1, column=1, value="М")
    lists.cell(row=2, column=1, value="Ж")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_LINES = (
    "Вуз: Московский технический университет связи и информатики",
    "Договор: ДГ-2025-017",
    "Программы: Python-разработчик, Инженер по тестированию",
    "Ответственный ИТ Школы: Петров Пётр Алексеевич",
    "Представитель вуза: см. карточку вуза",
)

_TEACHERS = (
    ("ФИО", "Должность", "Курс пройден"),
    [
        ("Гусева Анна Петровна", "Доцент кафедры", "да"),
        ("Тарасов Олег Юрьевич", "Старший преподаватель", "да"),
        ("Власова Ирина Николаевна", "Ассистент", "нет"),
    ],
)
_CURRICULUM = (
    ("Модуль", "Часы", "Форма контроля"),
    [
        ("Основы и инструменты", 36, "Зачёт"),
        ("Практикум на учебном стенде", 72, "Проект"),
        ("Промышленные практики", 36, "Экзамен"),
    ],
)

# файл -> (формат, заголовок документа, таблица для xlsx/xls)
ATTACHMENTS: dict[str, tuple[str, str, files.Table | None]] = {
    "meeting-protocol.pdf": ("pdf", "Протокол встречи с вузом", None),
    "signed-contract-scan.jpeg": ("jpeg", "Скан подписанного договора", None),
    "stand-diagram.png": ("png", "Схема учебного стенда", None),
    "contract-draft.docx": ("docx", "Проект договора", None),
    "contract-university-edition.doc": ("doc", "Договор — редакция вуза", None),
    "teachers.xlsx": ("xlsx", "Список обученных преподавателей", _TEACHERS),
    "curriculum.xls": ("xls", "Учебный план", _CURRICULUM),
    "materials.zip": ("zip", "Методические материалы", None),
    "deploy-log.gz": ("gz", "Журнал развёртывания стенда", None),
    "methodology.rar": ("rar", "Методические материалы", None),
}


def write(target: Path) -> list[Path]:
    written: list[Path] = []

    folder = target / "import"
    folder.mkdir(parents=True, exist_ok=True)
    for name, (title, build, kind) in IMPORTS.items():
        headers, rows = build()
        writer = files.xlsx if kind == "xlsx" else files.xls
        path = folder / name
        path.write_bytes(writer(title, headers, rows))
        written.append(path)

    path = folder / "09-lms-learners.xlsx"
    path.write_bytes(lms_workbook())
    written.append(path)

    folder = target / "integrations"
    folder.mkdir(parents=True, exist_ok=True)
    for name, payload in (
        ("site-applications.json", site_applications()),
        ("lms-learners.json", lms_learners()),
    ):
        path = folder / name
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=1) + "\n", "utf-8")
        written.append(path)

    folder = target / "attachments"
    folder.mkdir(parents=True, exist_ok=True)
    for name, (kind, title, table) in ATTACHMENTS.items():
        path = folder / name
        path.write_bytes(files.document(kind, title, _LINES, table))
        written.append(path)

    # Недопустимый формат: загрузка должна отказать с кодом file_type_not_allowed.
    path = folder / "not-allowed.txt"
    path.write_text("Текстовые файлы во вложения не принимаются.\n", encoding="utf-8")
    written.append(path)
    return written


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Файлы для ручной проверки загрузок")
    parser.add_argument("target", type=Path, help="куда положить файлы")
    for path in write(parser.parse_args().target):
        print(path)
