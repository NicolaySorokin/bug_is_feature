"""Небольшие, но настоящие файлы каждого формата из ТЗ.

Функциональное требование 3 перечисляет форматы вложений: png, jpeg, pdf,
zip, gzip, rar, doc, docx, xls, xlsx. Для каждого здесь есть генератор,
который собирает правдоподобный документ из заголовка и нескольких строк:
ими наполняются этапы демонстрационных договоров и каталог testdata.

Оговорки по трём форматам:

* doc - это RTF. Двоичный формат Word 97 без самого Word не собрать,
  а RTF с расширением .doc Word и LibreOffice открывают как обычный документ;
* rar - архив RAR 4 без сжатия. Формат хранения открыт, сжатие RAR - нет.
  Такой архив открывают WinRAR, 7-Zip и tar из Windows;
* xls собирается библиотекой xlwt - она стоит только в зависимостях
  разработки, поэтому в демоданных таблицы идут в xlsx.
"""

from __future__ import annotations

import gzip
import io
import struct
import zipfile
import zlib
from collections.abc import Sequence
from datetime import date, datetime
from xml.sax.saxutils import escape

from app.services.export import fonts

BRAND = "#7700FF"  # фиолетовый Ростелекома
FOOTER = "Демонстрационный документ: сформирован автоматически для показа системы"
# Дата внутри архивов фиксирована: одинаковые файлы собираются в одинаковые байты.
ARCHIVE_TIME = (2026, 9, 1, 10, 0, 0)

Table = tuple[Sequence[str], Sequence[Sequence[object]]]


# --- PDF ----------------------------------------------------------------------


def pdf(title: str, lines: Sequence[str]) -> bytes:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    regular, bold = fonts.ensure_fonts()
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4, invariant=True)
    page.setTitle(title)
    page.setAuthor("ИТ Школа Ростелекома")
    width, height = A4

    page.setFillColor(HexColor(BRAND))
    page.rect(0, height - 18 * mm, width, 18 * mm, stroke=0, fill=1)
    page.setFillColor(HexColor("#FFFFFF"))
    page.setFont(bold, 11)
    page.drawString(20 * mm, height - 11 * mm, "ИТ Школа Ростелекома")

    page.setFillColor(HexColor("#1F1F24"))
    y = height - 34 * mm
    for part in simpleSplit(title, bold, 16, width - 40 * mm):
        page.setFont(bold, 16)
        page.drawString(20 * mm, y, part)
        y -= 8 * mm

    y -= 4 * mm
    page.setFont(regular, 11)
    for line in lines:
        for part in simpleSplit(line, regular, 11, width - 40 * mm) or [""]:
            page.drawString(20 * mm, y, part)
            y -= 6.5 * mm

    page.setFillColor(HexColor("#8C8C96"))
    page.setFont(regular, 8)
    page.drawString(20 * mm, 12 * mm, FOOTER)
    page.showPage()
    page.save()
    return buffer.getvalue()


# --- Изображения --------------------------------------------------------------


def _font(size: int, bold: bool = False):  # noqa: ANN202 - тип из Pillow
    from PIL import ImageFont

    found = fonts.font_files()
    if found is None:
        return ImageFont.load_default()
    return ImageFont.truetype(str(found[1] if bold else found[0]), size)


def _wrap(text: str, font, width: int) -> list[str]:  # noqa: ANN001 - шрифт Pillow
    words, lines, current = text.split(), [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if font.getlength(candidate) <= width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    return [*lines, current] if current else lines


def image(title: str, lines: Sequence[str], kind: str = "png") -> bytes:
    """PNG - схема на белом фоне, JPEG - «скан» документа с печатью."""
    from PIL import Image, ImageDraw

    scan = kind in {"jpeg", "jpg"}
    size = (760, 980) if scan else (1000, 560)
    canvas = Image.new("RGB", size, (246, 243, 236) if scan else (255, 255, 255))
    draw = ImageDraw.Draw(canvas)
    head, body = _font(30, bold=True), _font(20)

    if not scan:
        draw.rectangle((0, 0, size[0], 64), fill=BRAND)
        draw.text((32, 18), "ИТ Школа Ростелекома", font=_font(22, bold=True), fill="white")

    y = 110 if scan else 96
    for part in _wrap(title, head, size[0] - 80):
        draw.text((40, y), part, font=head, fill=(31, 31, 36))
        y += 42
    y += 16
    for line in lines:
        for part in _wrap(line, body, size[0] - 80):
            draw.text((40, y), part, font=body, fill=(60, 60, 66))
            y += 30

    if scan:
        # Круглая печать и подпись - чтобы скан был похож на скан.
        cx, cy = size[0] - 190, size[1] - 190
        draw.ellipse((cx - 110, cy - 110, cx + 110, cy + 110), outline=(40, 70, 160), width=6)
        draw.text(
            (cx - 74, cy - 14), "ПОДПИСАНО", font=_font(26, bold=True), fill=(40, 70, 160)
        )
        draw.line((60, size[1] - 150, 330, size[1] - 175), fill=(20, 30, 90), width=4)
    else:
        # Три блока со стрелками - условная схема стенда.
        for index, label in enumerate(("Вуз", "Стенд", "ИТ Школа")):
            left = 60 + index * 320
            draw.rounded_rectangle(
                (left, size[1] - 150, left + 240, size[1] - 70),
                radius=14,
                outline=BRAND,
                width=4,
            )
            draw.text((left + 24, size[1] - 124), label, font=body, fill=BRAND)
            if index < 2:
                draw.line(
                    (left + 240, size[1] - 110, left + 320, size[1] - 110), fill=BRAND, width=4
                )

    buffer = io.BytesIO()
    if scan:
        canvas.save(buffer, format="JPEG", quality=72, optimize=True)
    else:
        canvas.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


# --- Документы Word -----------------------------------------------------------

_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/'
    'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
    '2006/relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)


def docx(title: str, paragraphs: Sequence[str]) -> bytes:
    """Минимальный пакет Office Open XML: больше Word для открытия не нужно."""
    body = [
        '<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/>'
        f'<w:sz w:val="32"/></w:rPr><w:t>{escape(title)}</w:t></w:r></w:p>'
    ]
    body += [
        f'<w:p><w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
        for text in paragraphs
    ]
    body.append(
        f'<w:p><w:r><w:rPr><w:i/><w:sz w:val="18"/></w:rPr><w:t>{escape(FOOTER)}</w:t>'
        "</w:r></w:p>"
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{''.join(body)}"
        '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="850" '
        'w:bottom="1134" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>'
        "</w:sectPr></w:body></w:document>"
    )
    return zip_archive(
        {
            "[Content_Types].xml": _CONTENT_TYPES.encode(),
            "_rels/.rels": _RELS.encode(),
            "word/document.xml": document.encode(),
        }
    )


def _rtf_text(text: str) -> str:
    out = []
    for char in text:
        code = ord(char)
        if char in "\\{}":
            out.append("\\" + char)
        elif code < 128:
            out.append(char)
        else:
            # \uN - знак Юникода; N - знаковое 16-битное число, за ним замена «?».
            out.append(f"\\u{code if code < 32768 else code - 65536}?")
    return "".join(out)


def doc(title: str, paragraphs: Sequence[str]) -> bytes:
    """Документ для расширения .doc: RTF, который Word открывает без вопросов."""
    parts = [
        r"{\rtf1\ansi\ansicpg1251\deff0{\fonttbl{\f0\fswiss\fcharset204 Arial;}}",
        r"\uc1\pard\qc\b\fs32 " + _rtf_text(title) + r"\b0\par",
        r"\pard\ql\fs24",
    ]
    parts += [_rtf_text(text) + r"\par" for text in paragraphs]
    parts.append(r"\i\fs18 " + _rtf_text(FOOTER) + r"\i0\par}")
    return "\n".join(parts).encode("ascii")


# --- Таблицы ------------------------------------------------------------------


def _sheet_title(title: str) -> str:
    """Имя листа Excel: до 31 знака и без запрещённых символов."""
    cleaned = "".join(" " if char in "\\/?*[]:" else char for char in title)
    return cleaned.strip()[:31] or "Лист1"


def xlsx(title: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(title)
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EFE6FF")
    for row in rows:
        sheet.append(list(row))
    for index, header in enumerate(headers, start=1):
        values = [str(header), *(str(row[index - 1] or "") for row in rows)]
        width = min(max(len(value) for value in values) + 3, 60)
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xls(title: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> bytes:
    """Старый формат Excel 97. Нужен xlwt из requirements-dev.txt."""
    try:
        import xlwt
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise RuntimeError("Для файлов .xls установите xlwt: pip install xlwt") from exc

    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet(_sheet_title(title))
    header_style = xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour lavender")
    date_style = xlwt.easyxf(num_format_str="DD.MM.YYYY")
    for column, header in enumerate(headers):
        sheet.write(0, column, header, header_style)
        values = [str(header), *(str(row[column] or "") for row in rows)]
        sheet.col(column).width = 256 * min(max(len(value) for value in values) + 3, 60)
    for index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            if isinstance(value, date):
                sheet.write(index, column, value, date_style)
            elif value is not None:
                sheet.write(index, column, value)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- Архивы -------------------------------------------------------------------


def zip_archive(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            info = zipfile.ZipInfo(name, date_time=ARCHIVE_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)
    return buffer.getvalue()


def gzip_file(name: str, data: bytes) -> bytes:
    """gzip хранит имя исходного файла в латинице - кириллицу туда не передаём."""
    buffer = io.BytesIO()
    with gzip.GzipFile(filename=name, mode="wb", fileobj=buffer, mtime=0) as archive:
        archive.write(data)
    return buffer.getvalue()


def _dos_time(moment: datetime) -> int:
    return (
        (moment.year - 1980) << 25
        | moment.month << 21
        | moment.day << 16
        | moment.hour << 11
        | moment.minute << 5
        | moment.second // 2
    )


def _rar_block(fields: bytes) -> bytes:
    """Заголовок блока RAR: младшие 16 бит CRC32 от всех полей после самого CRC."""
    return struct.pack("<H", zlib.crc32(fields) & 0xFFFF) + fields


def rar_archive(files: dict[str, bytes]) -> bytes:
    """Архив RAR 4 без сжатия (метод «хранение»). Имена файлов - латиницей."""
    out = bytearray(b"Rar!\x1a\x07\x00")  # сигнатура RAR 1.5-4.x
    # Заголовок архива: тип 0x73, флагов нет, размер 13, два зарезервированных поля.
    out += _rar_block(struct.pack("<BHHHI", 0x73, 0x0000, 13, 0, 0))

    stamp = _dos_time(datetime(*ARCHIVE_TIME))
    for name, data in files.items():
        encoded = name.encode("ascii")
        header = struct.pack(
            "<BHHIIBIIBBHI",
            0x74,  # тип блока: файл
            0x8000,  # за заголовком идут данные файла
            32 + len(encoded),  # размер заголовка вместе с CRC и именем
            len(data),  # размер в архиве
            len(data),  # исходный размер - без сжатия они совпадают
            2,  # создан в Windows
            zlib.crc32(data),
            stamp,
            20,  # для распаковки хватит RAR 2.0
            0x30,  # метод: хранение без сжатия
            len(encoded),
            0x20,  # атрибут «архивный»
        )
        out += _rar_block(header + encoded) + data

    out += _rar_block(struct.pack("<BHH", 0x7B, 0x4000, 7))  # конец архива
    return bytes(out)


# --- Выбор генератора по расширению -------------------------------------------


def document(kind: str, title: str, lines: Sequence[str], table: Table | None = None) -> bytes:
    """Файл нужного формата с заданным содержимым."""
    if kind == "pdf":
        return pdf(title, lines)
    if kind in {"png", "jpeg", "jpg"}:
        return image(title, lines, kind)
    if kind == "docx":
        return docx(title, lines)
    if kind == "doc":
        return doc(title, lines)
    if kind in {"xlsx", "xls"}:
        headers, rows = table or (("№", "Содержание"), list(enumerate(lines, start=1)))
        return xlsx(title, headers, rows) if kind == "xlsx" else xls(title, headers, rows)
    if kind == "zip":
        return zip_archive(
            {
                f"{_ascii_slug(title)}.pdf": pdf(title, lines),
                f"{_ascii_slug(title)}.docx": docx(title, lines),
            }
        )
    if kind in {"gz", "gzip"}:
        log = "\n".join(
            f"2026-09-01 10:{index:02d}:00 INFO {line}" for index, line in enumerate(lines)
        )
        return gzip_file("deploy.log", log.encode())
    if kind == "rar":
        return rar_archive(
            {
                "metodichka.pdf": pdf(title, lines),
                "README.txt": "\n".join(lines).encode("utf-8"),
            }
        )
    raise ValueError(f"Нет генератора для формата «{kind}»")


def _ascii_slug(text: str) -> str:
    from scripts.demo.people import translit

    slug = "".join(char if char.isalnum() else "-" for char in translit(text))
    return "-".join(part for part in slug.split("-") if part)[:40] or "document"
